# -*- coding: utf-8 -*-
"""
Build the Arabic NON-TECHNICAL walkthrough presentation:
  FLAME_Overview_AR.pptx
Same screenshots and structure as the English deck, with Arabic text and RTL layout.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image as PImage

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")

FLAME = (0xF0, 0x4A, 0x18)
FLAME_DEEP = (0xA7, 0x19, 0x14)
SKY = (0x34, 0xA9, 0xD6)
LEAF = (0x27, 0x95, 0x6A)
INK = (0x24, 0x1C, 0x1A)
MUTED = (0x6B, 0x5E, 0x5A)
CREAM = (0xFF, 0xFA, 0xF4)
AR_FONT = "Arial"  # widely available Arabic-capable font

def shot(name):
    p = os.path.join(SHOTS, name)
    return p if os.path.exists(p) else None

# section, title(ar), intro(ar), steps(ar)[], screenshot
SLIDES = [
    ("ترحيب", "مرحبًا بك في منصة FLAME",
     "نافذة بسيطة على خدمة FLAME التعليمية — تابع قصة الخدمة وأدِر كل شيء من مكان واحد.",
     ["افتح الموقع من أي جهاز.",
      "الزوّار يتصفّحون الصفحات العامة.",
      "الفريق يسجّل الدخول لإدارة البيانات."],
     "a01_home.png"),

    ("للزوّار", "التقديم للانضمام",
     "يمكن لأي شخص التقديم للخدمة مباشرة من الصفحة الرئيسية.",
     ["اضغط على زر «تقديم».",
      "تظهر استمارة بسيطة.",
      "املأها وأرسلها — يستلمها الفريق فورًا."],
     "a02_apply_form.png"),

    ("للزوّار", "قصة الخدمة (لوحة المعلومات)",
     "ملخّص حيّ وملوّن لما يحدث عبر المدارس.",
     ["افتح لوحة المعلومات.",
      "شاهد الإجماليات: المهام، المدارس، المكتمل، الجاري، لم يبدأ.",
      "مرّر لاستكشاف تقدّم كل مدرسة."],
     "a03_dashboard.png"),

    ("للزوّار", "اختيار المدة بسهولة",
     "اختر أي فترة بتقويم سهل — دون كتابة التواريخ.",
     ["اضغط على خانة التاريخ.",
      "استخدم التقويم لاختيار يوم، أو «اليوم» / «مسح».",
      "تتحدّث لوحة المعلومات حسب الفترة."],
     "a04_datepicker.png"),

    ("للزوّار", "استعراض المهام",
     "شاهد كل مهمة بحالتها وتفاصيلها.",
     ["افتح صفحة المهام.",
      "كل مهمة تُظهر المكان والزمان والتقدّم.",
      "الحالة واضحة: مكتمل، جارٍ، أو لم يبدأ."],
     "a05_missions.png"),

    ("للزوّار", "استكشاف خريطة العالم",
     "شاهد أين تصل الخدمة على خريطة تفاعلية.",
     ["افتح صفحة الخريطة.",
      "اضغط على فقاعة الدولة للتركيز عليها.",
      "اضغط على علامة المدرسة لرؤية مهامها وأفرادها."],
     "a06_map.png"),

    ("للزوّار", "تقويم أعياد الميلاد",
     "لا تفوّت أي مناسبة — مع نسخ رقم التواصل بلمسة واحدة.",
     ["افتح صفحة أعياد الميلاد.",
      "اختر يومًا لرؤية صاحب عيد الميلاد.",
      "اضغط على رقم الهاتف لنسخه."],
     "a07_birthdays.png"),

    ("للزوّار", "مكتبة المناهج",
     "اطّلع على مواد المناهج والقيم والتكوين المشتركة.",
     ["افتح صفحة المناهج.",
      "تنقّل بين تبويبات: المناهج، القيم، التكوين.",
      "افتح أي مستند لقراءته."],
     "a08_curriculum.png"),

    ("للفريق", "تسجيل دخول آمن",
     "يسجّل الفريق الدخول لإدارة كل شيء بأمان.",
     ["افتح صفحة تسجيل الدخول.",
      "أدخل بريدك وكلمة المرور.",
      "ترى فقط ما تسمح به صلاحيتك."],
     "a09_login.png"),

    ("للفريق", "الصفحة الرئيسية للإدارة",
     "كل ما يديره الفريق منظّم في تبويبات بسيطة.",
     ["بعد تسجيل الدخول تصل إلى منطقة الإدارة.",
      "استخدم التبويبات: البيانات، المهام، الأخبار، المستخدمون، النماذج.",
      "كل تبويب يحوي أقسامًا سهلة قابلة للطيّ."],
     "a10_admin_home.png"),

    ("للفريق", "إنشاء المهام ومتابعتها",
     "خطّط مهمة، اختر المدارس والأفراد، وتابع حالتها.",
     ["افتح تبويب المهام.",
      "أضف العنوان والتواريخ والحالة والمدارس المزارة.",
      "عدّل أي مهمة لاحقًا — ويظهر تقدّمها في لوحة المعلومات."],
     "a11_admin_missions.png"),

    ("للفريق", "إدارة مدرسة",
     "افتح مدرسة لإضافة السجلات والمهام والأفراد والمحتوى — في مكان واحد.",
     ["افتح البيانات ← المدارس ووسّع مدرسة.",
      "أضف سجلًّا أو مهمة (طبّي، زيارة مدير، إلخ).",
      "عيّن المديرين/المعلمين وحدّث الحالات."],
     "a12_admin_schools.png"),

    ("للفريق", "تصميم استمارة التقديم",
     "حدّد ما يُسأل عنه المتقدّمون — دون الحاجة لمبرمج.",
     ["افتح تبويب استمارة التقديم.",
      "أضف الحقول: نص، رقم، قائمة، زر تبديل، أو تقييم.",
      "تظهر التغييرات في الاستمارة العامة فورًا."],
     "a13_admin_formbuilder.png"),

    ("للفريق", "إدارة أعضاء الفريق",
     "يضيف المسؤول أعضاء الفريق ويحدّد صلاحيات كلٍّ منهم.",
     ["افتح تبويب المستخدمين (للمسؤول فقط).",
      "أضف عضوًا وعيّن له صلاحية أو أكثر.",
      "الصلاحيات تحدّد من يحرّر المحتوى أو يراجع أو يدير النماذج."],
     "a14_admin_users.png"),
]

SECTION_COLORS = {"ترحيب": FLAME_DEEP, "للزوّار": SKY, "للفريق": LEAF}

def rgb(t): return RGBColor(*t)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW = prs.slide_width; blank = prs.slide_layouts[6]

def set_rtl(paragraph):
    pPr = paragraph._pPr
    if pPr is None:
        pPr = paragraph._p.get_or_add_pPr()
    pPr.set('rtl', '1')

def bg(s, c): s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(c)
def box(s, l, t, w, h, c, rad=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(c); shp.line.fill.background(); shp.shadow.inherit = False
    return shp
def txt(s, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.RIGHT, font=AR_FONT, rtl=True):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    if rtl: set_rtl(p)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = font
    return tb

# Title slide
s = prs.slides.add_slide(blank); bg(s, FLAME_DEEP)
box(s, 0, Inches(2.4), SW, Inches(2.7), FLAME)
logo = os.path.join(HERE, "client", "public", "flame-logo.png")
if os.path.exists(logo):
    try:
        s.shapes.add_picture(logo, int(SW/2 - Inches(0.65)), Inches(0.7), height=Inches(1.3))
    except Exception: pass
txt(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.2), "منصة FLAME", 54, (255,255,255), True, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(3.8), Inches(11.7), Inches(0.8), "جولة مبسّطة", 28, (255,255,255), False, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(0.6), "ما يمكنك مشاهدته والقيام به — خطوة بخطوة", 16, (255,224,214), False, PP_ALIGN.CENTER)

def divider(title, color):
    s = prs.slides.add_slide(blank); bg(s, color)
    txt(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.2), title, 44, (255,255,255), True, PP_ALIGN.CENTER)

current_section = None
num = 0
for section, title, intro, steps, img in SLIDES:
    if section != current_section:
        divider(section, SECTION_COLORS.get(section, FLAME_DEEP))
        current_section = section
    num += 1
    s = prs.slides.add_slide(blank); bg(s, CREAM)
    accent = SECTION_COLORS.get(section, FLAME)
    box(s, 0, 0, SW, Inches(1.15), accent, rad=False)
    txt(s, Inches(1.0), Inches(0.16), Inches(11.3), Inches(0.5), title, 27, (255,255,255), True, PP_ALIGN.RIGHT)
    txt(s, Inches(1.0), Inches(0.78), Inches(11.3), Inches(0.32), intro, 13, (255,255,255), False, PP_ALIGN.RIGHT)

    # Steps card on the RIGHT (RTL) ; screenshot on the LEFT
    box(s, Inches(7.85), Inches(1.5), Inches(5.0), Inches(5.4), (255,255,255))
    txt(s, Inches(8.05), Inches(1.68), Inches(4.6), Inches(0.4), "الخطوات", 17, accent, True, PP_ALIGN.RIGHT)
    yy = Inches(2.4)
    for i, st in enumerate(steps, 1):
        tb = s.shapes.add_textbox(Inches(8.05), yy, Inches(4.6), Inches(1.0)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; set_rtl(p)
        r2 = p.add_run(); r2.text = st + "  "; r2.font.size = Pt(13.5); r2.font.color.rgb = rgb(INK); r2.font.name = AR_FONT
        r1 = p.add_run(); r1.text = f" {i}"; r1.font.bold = True; r1.font.size = Pt(15); r1.font.color.rgb = rgb(accent); r1.font.name = AR_FONT
        yy = Emu(int(yy) + Inches(0.85))

    sp = shot(img)
    if sp:
        iw, ih = PImage.open(sp).size
        maxw, maxh = Inches(6.9), Inches(5.3); ratio = min(maxw/iw, maxh/ih)
        w = Emu(int(iw*ratio)); h = Emu(int(ih*ratio))
        lft = Inches(0.5) + (maxw - w)//2
        box(s, lft-Emu(55000), Inches(1.55)-Emu(55000), w+Emu(110000), h+Emu(110000), (255,255,255))
        s.shapes.add_picture(sp, lft, Inches(1.55), width=w, height=h)

    txt(s, Inches(0.3), Inches(7.05), Inches(1.4), Inches(0.4), f"{num}", 12, MUTED, False, PP_ALIGN.LEFT, rtl=False)

# Closing
s = prs.slides.add_slide(blank); bg(s, FLAME_DEEP)
txt(s, Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.0), "هذه هي منصة FLAME", 46, (255,255,255), True, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.6), "بسيطة للزوّار. قويّة للفريق.", 20, (255,224,214), False, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(5.6), Inches(11.7), Inches(0.5), "خدمة FLAME · كنيسة St. Paul", 14, (255,220,210), False, PP_ALIGN.CENTER)

out = os.path.join(HERE, "FLAME_Overview_AR.pptx")
prs.save(out)
print("Saved", out, "with", len(prs.slides._sldIdLst), "slides")
