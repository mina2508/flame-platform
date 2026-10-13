# -*- coding: utf-8 -*-
"""
Build a NON-TECHNICAL walkthrough presentation:
  FLAME_Overview.pptx
One friendly slide per action, with simple steps and a screenshot.
No code, no APIs, no database jargon — just what you see and do.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image as PImage

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")

FLAME = (0xF0, 0x4A, 0x18)
FLAME_DEEP = (0xA7, 0x19, 0x14)
SKY = (0x34, 0xA9, 0xD6)
LEAF = (0x27, 0x95, 0x6A)
INK = (0x24, 0x1C, 0x1A)
MUTED = (0x6B, 0x5E, 0x5A)
CREAM = (0xFF, 0xFA, 0xF4)

def shot(name):
    p = os.path.join(SHOTS, name)
    return p if os.path.exists(p) else None

# section, title, friendly intro, steps[], screenshot
SLIDES = [
    ("WELCOME", "Welcome to FLAME",
     "A simple window into the FLAME educational missions — see the service story and manage everything in one place.",
     ["Open the website on any device.",
      "Visitors explore the public pages.",
      "The team signs in to manage the data."],
     "a01_home.png"),

    ("FOR VISITORS", "Apply to Join",
     "Anyone can apply to serve, right from the home page.",
     ["Click the Apply button.",
      "A simple form pops up.",
      "Fill it in and submit — the team receives it instantly."],
     "a02_apply_form.png"),

    ("FOR VISITORS", "The Missions Story (Dashboard)",
     "A live, colourful summary of what is happening across the schools.",
     ["Open the Dashboard.",
      "See totals: missions, schools, completed, pending, not started.",
      "Scroll to explore each school's progress."],
     "a03_dashboard.png"),

    ("FOR VISITORS", "Pick a Date Range Easily",
     "Choose any period with a friendly calendar — no typing dates.",
     ["Click a date field.",
      "Use the calendar to pick a day, or Today / Clear.",
      "The whole dashboard updates to that period."],
     "a04_datepicker.png"),

    ("FOR VISITORS", "Browse Missions",
     "See every mission with its status and details.",
     ["Open the Missions page.",
      "Each mission shows where, when and its progress.",
      "Status is clearly flagged: completed, in progress, or not started."],
     "a05_missions.png"),

    ("FOR VISITORS", "Explore the World Map",
     "See where the service reaches on an interactive map.",
     ["Open the Map page.",
      "Click a country bubble to focus on it.",
      "Click a school pin to see its missions and people."],
     "a06_map.png"),

    ("FOR VISITORS", "Birthday Calendar",
     "Never miss a celebration — with one-tap contact copying.",
     ["Open the Birthdays page.",
      "Pick a day to see whose birthday it is.",
      "Tap the phone chip to copy the contact number."],
     "a07_birthdays.png"),

    ("FOR VISITORS", "Curriculum Library",
     "Read the shared curriculum, values and formation materials.",
     ["Open the Curriculum page.",
      "Switch between Curriculum, Values and Formation tabs.",
      "Open any document to read it."],
     "a08_curriculum.png"),

    ("FOR THE TEAM", "Secure Sign-In",
     "The team signs in to manage everything safely.",
     ["Open the Login page.",
      "Enter your email and password.",
      "You only see what your role allows."],
     "a09_login.png"),

    ("FOR THE TEAM", "The Admin Home",
     "Everything the team manages, organised into simple tabs.",
     ["After signing in, you land on the admin area.",
      "Use the tabs: Database, Missions, News, Users, Forms.",
      "Each tab holds friendly, collapsible sections."],
     "a10_admin_home.png"),

    ("FOR THE TEAM", "Create & Track Missions",
     "Plan a mission, pick schools and people, and track its status.",
     ["Open the Missions tab.",
      "Add a title, dates, status and the schools visited.",
      "Edit any mission later — its progress shows on the dashboard."],
     "a11_admin_missions.png"),

    ("FOR THE TEAM", "Manage a School",
     "Open a school to add logs, tasks, people and content — all in one place.",
     ["Open Database → Schools and expand a school.",
      "Add a log or task (medical, principal visit, etc.).",
      "Assign principals/teachers and update statuses."],
     "a12_admin_schools.png"),

    ("FOR THE TEAM", "Design the Apply Form",
     "Decide what applicants are asked — no developer needed.",
     ["Open the Application Form tab.",
      "Add fields: text, number, dropdown, toggle or rating.",
      "Changes appear on the public Apply form instantly."],
     "a13_admin_formbuilder.png"),

    ("FOR THE TEAM", "Manage Team Members",
     "Admins add team members and choose what each can do.",
     ["Open the Users tab (admin only).",
      "Add a member and assign one or more roles.",
      "Roles control who edits content, reviews, or manages forms."],
     "a14_admin_users.png"),
]

def rgb(t): return RGBColor(*t)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW = prs.slide_width; blank = prs.slide_layouts[6]

def bg(s, c): s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(c)
def box(s, l, t, w, h, c, rad=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(c); shp.line.fill.background(); shp.shadow.inherit = False
    return shp
def txt(s, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Segoe UI"):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = font
    return tb

SECTION_COLORS = {"WELCOME": FLAME_DEEP, "FOR VISITORS": SKY, "FOR THE TEAM": LEAF}

# Title slide
s = prs.slides.add_slide(blank); bg(s, FLAME_DEEP)
box(s, 0, Inches(2.4), SW, Inches(2.7), FLAME)
logo = os.path.join(HERE, "client", "public", "flame-logo.png")
if os.path.exists(logo):
    try:
        iw, ih = PImage.open(logo).size; ratio = Inches(1.3)/ih
        s.shapes.add_picture(logo, int(SW/2 - iw*ratio/2), Inches(0.7), height=Inches(1.3))
    except Exception: pass
txt(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.2), "FLAME Platform", 54, (255,255,255), True, PP_ALIGN.CENTER, "Georgia")
txt(s, Inches(0.8), Inches(3.8), Inches(11.7), Inches(0.8), "A Simple Walkthrough", 26, (255,255,255), False, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(0.6), "What you can see and do — step by step", 16, (255,224,214), False, PP_ALIGN.CENTER)

# Section divider helper
def divider(title, color):
    s = prs.slides.add_slide(blank); bg(s, color)
    txt(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.2), title, 44, (255,255,255), True, PP_ALIGN.CENTER, "Georgia")

current_section = None
num = 0
for section, title, intro, steps, img in SLIDES:
    if section != current_section:
        divider(section, SECTION_COLORS.get(section, FLAME_DEEP))
        current_section = section
    num += 1
    s = prs.slides.add_slide(blank); bg(s, CREAM)
    accent = SECTION_COLORS.get(section, FLAME)
    box(s, 0, 0, SW, Inches(1.15), accent, rad=False)
    txt(s, Inches(0.6), Inches(0.16), Inches(11), Inches(0.5), title, 27, (255,255,255), True, font="Georgia")
    txt(s, Inches(0.62), Inches(0.78), Inches(12), Inches(0.32), intro, 13, (255,255,255))

    # Steps card (left)
    box(s, Inches(0.5), Inches(1.5), Inches(5.0), Inches(5.4), (255,255,255))
    txt(s, Inches(0.72), Inches(1.68), Inches(4.6), Inches(0.4), "Steps", 17, accent, True)
    yy = Inches(2.4)
    for i, st in enumerate(steps, 1):
        tb = s.shapes.add_textbox(Inches(0.72), yy, Inches(4.6), Inches(1.0)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = f"{i}  "; r1.font.bold = True; r1.font.size = Pt(15); r1.font.color.rgb = rgb(accent)
        r2 = p.add_run(); r2.text = st; r2.font.size = Pt(13.5); r2.font.color.rgb = rgb(INK)
        yy = Emu(int(yy) + Inches(0.85))

    # Screenshot (right)
    sp = shot(img)
    if sp:
        iw, ih = PImage.open(sp).size
        maxw, maxh = Inches(6.9), Inches(5.3); ratio = min(maxw/iw, maxh/ih)
        w = Emu(int(iw*ratio)); h = Emu(int(ih*ratio))
        lft = Inches(5.9) + (maxw - w)//2
        box(s, lft-Emu(55000), Inches(1.55)-Emu(55000), w+Emu(110000), h+Emu(110000), (255,255,255))
        s.shapes.add_picture(sp, lft, Inches(1.55), width=w, height=h)

    txt(s, Inches(11.7), Inches(7.05), Inches(1.4), Inches(0.4), f"{num}", 12, MUTED, False, PP_ALIGN.RIGHT)

# Closing
s = prs.slides.add_slide(blank); bg(s, FLAME_DEEP)
txt(s, Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.0), "That's FLAME", 46, (255,255,255), True, PP_ALIGN.CENTER, "Georgia")
txt(s, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.6), "Simple for visitors. Powerful for the team.", 20, (255,224,214), False, PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(5.6), Inches(11.7), Inches(0.5), "FLAME Function · St. Paul Coptic Service", 14, (255,220,210), False, PP_ALIGN.CENTER)

out = os.path.join(HERE, "FLAME_Overview.pptx")
prs.save(out)
print("Saved", out, "with", len(prs.slides._sldIdLst), "slides")
